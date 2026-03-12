#!/usr/bin/env python3
"""
生成水利厅售前演示PPT
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 定义配色方案
TITLE_COLOR = RGBColor(0, 82, 147)  # 深蓝色
ACCENT_COLOR = RGBColor(0, 120, 212)  # 浅蓝色
TEXT_COLOR = RGBColor(51, 51, 51)  # 深灰色
WHITE = RGBColor(255, 255, 255)

def add_title_slide(prs, title, subtitle=""):
    """添加标题页"""
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)

    # 背景形状
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = TITLE_COLOR
    shape.line.fill.background()

    # 主标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    # 副标题
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(1))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    return slide

def add_content_slide(prs, title, bullets):
    """添加内容页"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # 顶部蓝色条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TITLE_COLOR
    bar.line.fill.background()

    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # 内容区域
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(12)

    return slide

def add_two_column_slide(prs, title, left_title, left_bullets, right_title, right_bullets):
    """添加双栏内容页"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # 顶部蓝色条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TITLE_COLOR
    bar.line.fill.background()

    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # 左侧标题
    left_title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.5), Inches(0.5))
    tf = left_title_box.text_frame
    p = tf.paragraphs[0]
    p.text = left_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    # 左侧内容
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.1), Inches(5.5), Inches(4.5))
    tf = left_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(left_bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(8)

    # 右侧标题
    right_title_box = slide.shapes.add_textbox(Inches(7), Inches(1.5), Inches(5.5), Inches(0.5))
    tf = right_title_box.text_frame
    p = tf.paragraphs[0]
    p.text = right_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    # 右侧内容
    right_box = slide.shapes.add_textbox(Inches(7), Inches(2.1), Inches(5.5), Inches(4.5))
    tf = right_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(right_bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(8)

    return slide

# ============ 开始生成幻灯片 ============

# 1. 封面
add_title_slide(prs,
    "智慧水利视觉智能平台",
    "一站式AI解决方案 | 让水利管理更智能")

# 2. 目录
add_content_slide(prs, "目录", [
    "客户需求分析",
    "解决方案概述",
    "场景一：工程施工安全监管",
    "场景二：水利设施安全巡查",
    "场景三：水质监测与预警",
    "产品核心优势",
    "部署方案与服务保障"
])

# 3. 客户需求
add_two_column_slide(prs, "客户需求分析",
    "业务背景", [
        "水利工程施工现场安全监管",
        "大坝/河道/水库安全巡查",
        "水质实时监测与污染预警",
        "人工巡检效率低、覆盖有限",
        "问题发现滞后，应急响应慢"
    ],
    "客户特点", [
        "无技术团队，需零代码工具",
        "倾向本地化部署，数据安全",
        "需全新建设完整方案",
        "有AI专项预算"
    ])

# 4. 解决方案概述
add_content_slide(prs, "解决方案概述", [
    "一站式视觉智能平台：从数据采集到AI部署全流程",
    "无需编程基础，Web界面全流程操作",
    "私有化部署，数据完全留存本地",
    "支持国产操作系统与芯片",
    "6-9周完成从0到上线"
])

# 5. 场景一
add_two_column_slide(prs, "场景一：工程施工安全监管",
    "业务需求", [
        "施工人员是否佩戴安全帽？",
        "是否有人员进入危险区域？",
        "施工现场是否存在违规行为？",
        "工程车辆进出是否合规？"
    ],
    "我们的方案", [
        "安全帽/反光衣检测",
        "禁区入侵实时预警",
        "施工车辆进出统计",
        "人员活动热力图分析",
        "支持多路视频同时分析"
    ])

# 6. 场景二
add_two_column_slide(prs, "场景二：水利设施安全巡查",
    "业务需求", [
        "大坝表面是否有裂缝、破损？",
        "河道护坡是否完好？",
        "水库闸门是否正常开启？",
        "是否有漂浮物堆积？"
    ],
    "我们的方案", [
        "图像缺陷自动识别",
        "对比历史影像变化检测",
        "河道漂浮物/非法建筑检测",
        "SAM辅助快速标注",
        "支持无人机/机器人图像"
    ])

# 7. 场景三
add_two_column_slide(prs, "场景三：水质监测与预警",
    "业务需求", [
        "水体颜色异常是否意味着污染？",
        "是否有偷排暗排行为？",
        "饮用水源地是否有威胁？",
        "蓝藻水华能否提前预警？"
    ],
    "我们的方案", [
        "水体颜色异常分析",
        "水面温度热力图预警",
        "漂浮物/蓝藻/油污检测",
        "实时监测，预警响应<1分钟",
        "可对接水质传感器数据"
    ])

# 8. 核心优势
add_content_slide(prs, "产品核心优势", [
    "零技术门槛：无需编码，点点鼠标完成全流程",
    "私有化安全部署：数据不出内网，国产化适配",
    "完整功能矩阵：标注/训练/推理/管理一体化",
    "水利行业定制：预置水利模板、场景预训练模型"
])

# 9. 量化价值
add_content_slide(prs, "方案量化价值", [
    "巡检效率提升 5-10 倍",
    "人力巡检成本降低 70% 以上",
    "问题发现到预警 < 1 分钟",
    "7×24小时不间断智能监控",
    "积累水利数据资产，支撑长期决策"
])

# 10. 部署方案
add_two_column_slide(prs, "部署方案",
    "硬件建议", [
        "基础版：CPU服务器 × 1",
        "标准版：GPU(RTX4080/4090) × 1",
        "增强版：GPU × 2",
        "旗舰版：GPU集群"
    ],
    "实施周期", [
        "需求确认：1周",
        "定制开发：2-4周",
        "部署实施：1周",
        "模型训练：1-2周",
        "试运行：1周",
        "合计：6-9周"
    ])

# 11. 服务保障
add_content_slide(prs, "服务保障", [
    "1年免费质保，7×24小时响应",
    "每年2次免费现场巡检",
    "终身免费版本升级",
    "包含2次现场培训服务",
    "可选增值服务：驻场支持"
])

# 12. 结束页
add_title_slide(prs,
    "感谢聆听",
    "让我们携手推动智慧水利建设")

# 保存文件
output_path = "/root/wuyu/YOLO-/docs/plans/智慧水利售前演示.pptx"
prs.save(output_path)
print(f"PPT已生成: {output_path}")
